#!/usr/bin/env python3
"""Build the student handout (.docx) for the Capstone topic catalogue.

    single source of truth : projects/capstone/topic-catalogue.md
    generated artefact     : projects/capstone/topic-catalogue.docx

Why a script rather than a one-off conversion: the catalogue is re-generated
every term (seeds move, licences get re-checked, week numbers shift), and a
handout that silently drifts from the Markdown is worse than no handout.  So
this script re-derives the .docx from the .md and then *asserts* that nothing
was dropped on the way through.

Pipeline
    Markdown --(python-markdown + BeautifulSoup)--> HTML --(tencent-docx
    bundled engine)--> DOCX

The HTML -> DOCX step calls the tencent-docx plugin's bundled engine with the
*managed arm64* interpreter.  The plugin's own venv is x86_64 and cannot run on
this machine (no Rosetta), so we invoke the engine module directly:

    cd <plugin>/skills/html-to-docx/scripts && <python> -m html_to_docx convert ...

Requirements (already present in the managed venv):
    markdown, beautifulsoup4, lxml, python-docx

Usage, from the repo root:

    /Users/renzheng/.workbuddy/binaries/python/envs/default/bin/python \
        scripts/build_topic_handout.py

Exit code 0 = built and verified; 1 = verification failed (the .docx is left on
disk for inspection).
"""

from __future__ import annotations

import re
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
MD_PATH = ROOT / "projects" / "capstone" / "topic-catalogue.md"
OUT_PATH = ROOT / "projects" / "capstone" / "topic-catalogue.docx"

REPO_URL = "https://github.com/klausren/ai-application-development"
BRANCH = "main"

# ---------------------------------------------------------------- the seeds
# Kept here so a dropped card fails the build instead of shipping quietly.
SEED_IDS = (
    [f"V{i}" for i in range(1, 5)]
    + [f"T{i}" for i in range(1, 5)]
    + [f"C{i}" for i in range(1, 5)]
    + [f"D{i}" for i in range(1, 6)]
    + [f"A{i}" for i in range(1, 5)]
)

# --------------------------------------------------------------- the layout
# The engine reads column widths from `width: N%` on the FIRST row's cells and
# rewrites <w:gridCol> proportionally.  Keyed by (kind, column count).
COLUMN_WIDTHS: dict[str, list[int]] = {
    "card":     [19, 81],                        # the 21 seed cards, 2 cols
    "meta":     [22, 78],                        # the who/when/what block
    "dial":     [22, 78],                        # §0 rule ② dials
    "quickpick": [4, 8, 17, 12, 6, 7, 8, 38],    # §2, 8 cols (the wide one)
    "week":     [6, 21, 24, 24, 25],             # §3 track -> spine
    "risk":     [26, 42, 32],                    # §5 high-risk topics
    "data":     [19, 19, 9, 25, 18, 10],         # §6 dataset leads
    "compute":  [12, 88],                        # §1 compute badges
}


# ------------------------------------------------------------- the headings
# One spec, two consumers: `css()` renders it into the stylesheet, and
# `normalize_headings()` enforces it on the converted .docx.
#
# The second consumer is not redundant.  The engine maps <h1>..<h4> onto Word's
# built-in Heading styles, whose own definitions leak through: measured on this
# machine, the built-in Heading 4 renders the seed-card titles in blue italics
# with an underline, and a direct run colour does not override it.  So after
# conversion we write the intended formatting explicitly onto every heading run.
#
# level -> (point size, bold, hex colour, alignment)
HEADING_SPEC: dict[int, tuple[float, bool, str, str]] = {
    1: (21, True, "1F4E79", "center"),   # document title
    2: (14, True, "1F4E79", "left"),     # §0 .. §7
    3: (12, True, "1F4E79", "left"),     # Track V / T / C / D / A
    4: (11, True, "1A1A1A", "left"),     # the 21 seed cards
}
HEADING_FONT = "微软雅黑"


TEXT_BLOCK_TAGS = ("p", "li", "td", "th", "h1", "h2", "h3", "h4")
INLINE_TAGS = ("em", "strong", "code", "a", "span", "sub", "sup")
# Direct children of these are separate blocks, so a newline between them is
# inter-block spacing and must stay a newline.
BLOCK_CONTAINERS = {"body", "div", "blockquote", "section", "ul", "ol", "table",
                    "tr", "thead", "tbody"}
LATIN = set("abcdefghijklmnopqrstuvwxyzABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789")


def css() -> str:
    """Stylesheet.  Only properties the engine actually maps are used.

    Caveat worth remembering: the engine's `apply_font_family` takes the FIRST
    font in the list and writes it to w:ascii, w:hAnsi *and* w:eastAsia.  So a
    list like `"Times New Roman", "宋体"` would give Chinese text Times New
    Roman.  Every family below is a single dual-coverage face on purpose.
    """
    base = """
    @page {
        @bottom-left  { content: "52015CC3BV · Capstone Topic Catalogue · 选题参考库"; }
        @bottom-right { content: counter(page) " / " counter(pages); }
    }
    body { font-family: "宋体"; font-size: 10.5pt; color: #1a1a1a; line-height: 1.45; }

    p  { margin-top: 3pt; margin-bottom: 3pt; }

    ul, ol { margin-top: 3pt; margin-bottom: 3pt; }
    li { margin-bottom: 2pt; }

    blockquote { font-size: 9.5pt; color: #444444; background-color: #F4F6F8;
                 border-left: 2.5pt solid #9DB8D2; margin-top: 6pt; margin-bottom: 6pt; }

    table { width: 100%; font-size: 8.5pt; }
    th { background-color: #1F4E79; color: #FFFFFF; font-size: 8.5pt;
         padding: 3pt; border: 0.5pt solid #A9B7C6; }
    td { font-size: 8.5pt; padding: 3pt; border: 0.5pt solid #C6CDD5; }

    code { font-family: "Consolas"; font-size: 8.5pt; color: #A33A1E; }
    hr { border-bottom: 0.75pt solid #D0D5DA; }

    .subtitle { font-family: "微软雅黑"; font-size: 10.5pt; color: #44546A;
                text-align: center; margin-top: 2pt; margin-bottom: 0; }
    .subtitle2 { font-size: 9pt; color: #7A7A7A;
                 text-align: center; margin-top: 0; margin-bottom: 10pt; }
    """

    # Heading rules are generated from HEADING_SPEC so the stylesheet and the
    # post-conversion enforcement can never disagree.
    spacing = {
        1: "margin-bottom: 2pt;",
        2: "border-bottom: 1.5pt solid #1F4E79; margin-top: 16pt; margin-bottom: 6pt;",
        3: "margin-top: 13pt; margin-bottom: 4pt;",
        4: "background-color: #E8F0F8; margin-top: 11pt; margin-bottom: 3pt;",
    }
    rules = [
        f'h{level} {{ font-family: "{HEADING_FONT}"; font-size: {size}pt;'
        f' color: #{color}; font-weight: {"bold" if bold else "normal"};'
        f' text-align: {align}; {spacing[level]} }}'
        for level, (size, bold, color, align) in HEADING_SPEC.items()
    ]
    return base + "\n    " + "\n    ".join(rules) + "\n    "


def engine_dir() -> Path:
    """Locate the tencent-docx bundled html_to_docx package (highest version)."""
    roots = [
        Path.home() / ".workbuddy" / "plugins" / "cache" / "workbuddy-builtin" / "tencent-docx",
        Path.home() / ".workbuddy" / "plugins" / "marketplaces" / "workbuddy-builtin"
        / "builtin-plugins" / "tencent-docx",
    ]
    found: list[tuple[tuple[int, ...], Path]] = []
    for root in roots:
        for scripts in root.glob("*/skills/html-to-docx/scripts") if root.name == "tencent-docx" else \
                       [root / "skills" / "html-to-docx" / "scripts"]:
            if (scripts / "html_to_docx").is_dir():
                ver = tuple(int(n) for n in re.findall(r"\d+", scripts.parts[-4])[:3]) or (0,)
                found.append((ver, scripts))
    if not found:
        raise SystemExit(
            "Could not find the tencent-docx html_to_docx engine.\n"
            "Looked under ~/.workbuddy/plugins/{cache,marketplaces}/.../tencent-docx"
        )
    return sorted(found)[-1][1]


# Print substitutions applied to the Markdown text before rendering.
#
# Why: 💻 (U+1F4BB) and 🖥️ (U+1F5A5) have no text-presentation glyph in the
# common CJK office fonts, so a reader without a colour-emoji font gets an
# unreadable solid block -- measured on this machine, see the repo notes.  The
# other five symbols the catalogue uses (✅ ⚠️ ❌ 🟢 ★) render correctly and are
# left alone.  Only the *handout* is substituted; the Markdown keeps its emoji.
PRINT_SUBSTITUTIONS = [
    ("💻🖥️", "CPU/GPU"),
    ("💻", "CPU"),
    ("🖥️", "GPU"),
    ("🖥", "GPU"),
]


def preprocess_md(text: str) -> str:
    """Substitutions applied before the Markdown renderer sees it."""
    text = re.sub(r"^(\s*)- \[ \] ", r"\1- ☐ ", text, flags=re.M)
    text = re.sub(r"^(\s*)- \[[xX]\] ", r"\1- ☑ ", text, flags=re.M)
    for src, dst in PRINT_SUBSTITUTIONS:
        text = text.replace(src, dst)
    return text


def table_kind(ncols: int, first_header: str) -> str:
    """Pick a width profile.  Signatures are the catalogue's own headers."""
    first_header = first_header.strip()
    if ncols == 2 and not first_header:
        return "card"
    if ncols == 2 and first_header.lower().startswith(("who", "dial", "budget")):
        return "meta" if first_header.lower().startswith("who") else "dial"
    if ncols == 2:
        return "meta"
    if ncols == 8:
        return "quickpick"
    if ncols == 5:
        return "week"
    if ncols == 3:
        return "risk"
    if ncols == 6:
        return "data"
    return "meta"


def md_to_html(md_text: str) -> tuple[str, dict]:
    """Markdown -> the styled HTML that the engine will convert."""
    import markdown
    from bs4 import BeautifulSoup

    body_md = preprocess_md(md_text)
    raw = markdown.markdown(
        body_md, extensions=["tables", "sane_lists", "fenced_code"]
    )
    soup = BeautifulSoup(raw, "lxml")

    stats = {"tables": 0, "empty_headers_dropped": 0, "links_rewritten": 0,
             "blockquotes_inlined": 0, "soft_breaks_joined": 0, "glue_risks": []}

    # 0. The engine drops a text node's leading whitespace when that node
    #    begins with a newline, which glues "<strong>…menu.</strong>\n21 seeds"
    #    into "menu.21 seeds".  HTML says a newline in text is only a space, so
    #    normalise it here.
    #
    #    Whitespace-only nodes count too *inside a text block* -- that is the
    #    newline between two inline runs, which is what glued
    #    "…School of Software, DNUI*\n*This catalogue…".  Directly under a block
    #    container the newline merely separates blocks, and turning it into a
    #    space would litter the page with blank paragraphs, so it stays.
    for node in soup.find_all(string=True):
        raw = str(node)
        if "\n" not in raw:
            continue
        if not raw.strip() and node.parent.name in BLOCK_CONTAINERS:
            continue
        normalised = re.sub(r"[ \t]*\n[ \t]*", " ", raw)
        if normalised != raw:
            node.replace_with(normalised)
            stats["soft_breaks_joined"] += 1

    # 0a. Regression guard for exactly that bug: after step 0 no inline element
    #     may still be followed directly by text that would fuse with it.
    #     Only a Latin/digit-to-Latin/digit boundary counts -- "**必须**拿到"
    #     and "…条款</code>。" are correct Chinese setting, not glue.
    for block in soup.find_all(TEXT_BLOCK_TAGS):
        for child in block.children:
            if getattr(child, "name", None) not in INLINE_TAGS:
                continue
            nxt = child.next_sibling
            if nxt is None or getattr(nxt, "name", None):
                continue
            tail = str(nxt)
            head = child.get_text().rstrip()
            if not tail or not head or tail[0].isspace():
                continue
            last, first = head[-1], tail[0]
            if (last in LATIN and first in LATIN) or (last == "." and first in LATIN):
                stats["glue_risks"].append(
                    f"{head[-24:]!r} would fuse with {tail[:24]!r}"
                )

    # 0b. Markdown wraps block quotes as <blockquote><p>…</p></blockquote>, but
    #     the engine's blockquote handler only reads a direct text child -- so
    #     the "read this / watch out" asides lost their tint and left rule.
    #     Inline the style onto the paragraph, which keeps the bold runs.
    for bq in soup.find_all("blockquote"):
        ps = bq.find_all("p")
        if ps:
            for p in ps:
                p["style"] = (p.get("style", "") + "background-color: #F4F6F8;"
                              " border-left: 2.5pt solid #9DB8D2;"
                              " margin-top: 6pt; margin-bottom: 6pt;")
            bq.unwrap()
            stats["blockquotes_inlined"] += 1

    # 1. Half the catalogue's tables are label/value grids written as `| | |`,
    #    i.e. with an empty header row.  Left alone, that ships a blank row.
    for thead in soup.find_all("thead"):
        if not thead.get_text(strip=True):
            thead.decompose()
            stats["empty_headers_dropped"] += 1

    # 2. Relative links are useless in a handout the student opens off the repo
    #    or off a printout -- point them at the published file.
    for a in soup.find_all("a", href=True):
        href = a["href"]
        if href.startswith(("http://", "https://", "#", "mailto:")):
            continue
        path, _, anchor = href.partition("#")
        target = (MD_PATH.parent / path).resolve()
        try:
            rel = target.relative_to(ROOT).as_posix()
        except ValueError:
            continue
        a["href"] = f"{REPO_URL}/blob/{BRANCH}/{rel}" + (f"#{anchor}" if anchor else "")
        stats["links_rewritten"] += 1

    # 3. Column widths have to be inline on the first row, and the engine reads
    #    them there -- CSS in <style> would not reach it.
    for table in soup.find_all("table"):
        stats["tables"] += 1
        rows = table.find_all("tr")
        if not rows:
            continue
        cells = rows[0].find_all(["th", "td"])
        ncols = len(cells)
        kind = table_kind(ncols, cells[0].get_text(strip=True) if cells else "")
        widths = COLUMN_WIDTHS.get(kind)
        if widths and len(widths) == ncols:
            for cell, width in zip(cells, widths):
                cell["style"] = f"width: {width}%;"

    content = str(soup)

    title = (
        '<h1>Capstone Topic Catalogue · 选题参考库</h1>'
        '<p class="subtitle">52015CC3BV · AI Application Development · '
        'Semester 1, 2026–2027 · 留软工 24401</p>'
        '<p class="subtitle2">School of Software, Dalian Neusoft University of Information'
        ' · Instructor: Zheng Ren 任政</p>'
    )
    # The Markdown already carries its own H1; the styled one replaces it.
    content = re.sub(r"<h1>.*?</h1>", "", content, count=1, flags=re.S)

    doc = (
        "<!DOCTYPE html><html><head><meta charset='utf-8'>"
        f"<style>{css()}</style></head><body>{title}{content}</body></html>"
    )
    return doc, stats


def run_engine(html_path: Path, out_path: Path) -> None:
    scripts = engine_dir()
    cmd = [
        sys.executable, "-m", "html_to_docx", "convert", str(html_path),
        "-o", str(out_path),
        "--page-size", "A4", "--orientation", "portrait",
        "--margin-top", "1.9", "--margin-bottom", "1.8",
        "--margin-left", "1.6", "--margin-right", "1.6",
    ]
    proc = subprocess.run(cmd, cwd=str(scripts), capture_output=True, text=True)
    if proc.returncode != 0:
        raise SystemExit(f"engine failed (exit {proc.returncode}):\n{proc.stderr[-2000:]}")
    print(f"  engine: {proc.stdout.strip()}")


def normalize_headings(path: Path) -> int:
    """Write HEADING_SPEC onto every heading run of the converted document.

    Word's built-in Heading styles carry their own definitions, and the engine's
    run-level CSS does not always win against them (Heading 4 in particular
    renders the seed-card titles in blue italics with an underline).  Setting
    the properties explicitly here -- including `italic=False` / `underline=False`,
    which emit `w:i val="0"` / `w:u val="none"` -- makes the result deterministic
    on any renderer.
    """
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    doc = Document(str(path))
    touched = 0
    for para in doc.paragraphs:
        match = re.match(r"Heading (\d+)$", para.style.name)
        if not match:
            continue
        spec = HEADING_SPEC.get(int(match.group(1)))
        if not spec:
            continue
        size, bold, colour, align = spec
        para.alignment = (WD_ALIGN_PARAGRAPH.CENTER if align == "center"
                          else WD_ALIGN_PARAGRAPH.LEFT)
        for run in para.runs:
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.italic = False
            run.font.underline = False
            run.font.color.rgb = RGBColor.from_string(colour)
            rpr = run._r.get_or_add_rPr()
            rfonts = rpr.find(qn("w:rFonts"))
            if rfonts is None:
                rfonts = OxmlElement("w:rFonts")
                rpr.insert(0, rfonts)
            for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                rfonts.set(qn(attr), HEADING_FONT)
        touched += 1
    doc.save(str(path))
    return touched


def docx_text(path: Path) -> tuple[str, list, list]:
    from docx import Document

    doc = Document(str(path))
    heads: list[str] = []
    chunks: list[str] = []
    for p in doc.paragraphs:
        chunks.append(p.text)
        if p.style.name.startswith("Heading") and p.text.strip():
            heads.append(p.text.strip())
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                chunks.append(c.text)
    return "\n".join(chunks), heads, doc.tables


def verify(md_text: str, docx_path: Path, stats: dict) -> list[str]:
    """Same-source assertions: the handout must carry the whole catalogue."""
    problems: list[str] = []
    blob, heads, tables = docx_text(docx_path)

    # (a) every seed card survives, matched on its full heading
    card_heads = re.findall(r"^####\s+([A-Z]\d)\s+·\s+(.+)$", md_text, flags=re.M)
    if len(card_heads) != len(SEED_IDS):
        problems.append(
            f"source has {len(card_heads)} seed cards, expected {len(SEED_IDS)}"
        )
    for sid, title in card_heads:
        needle = f"{sid} · {title.strip()}"
        if needle not in blob:
            problems.append(f"handout is missing card: {needle[:60]}")

    # (b) table count survives
    if len(tables) != stats["tables"]:
        problems.append(
            f"handout has {len(tables)} tables, HTML had {stats['tables']}"
        )

    # (c) no Markdown syntax leaked through as literal text
    for token, label in (("**", "bold markers"), ("`", "backticks"), ("|---", "table rules")):
        if token in blob:
            problems.append(f"literal {label} ({token!r}) survived into the .docx")

    # (d) the two emoji that render as unreadable blocks must have been swapped
    for bad, label in (("\U0001f4bb", "💻"), ("\U0001f5a5", "🖥")):
        if bad in blob:
            problems.append(f"unrenderable {label} survived into the .docx")

    # (e) regression guard for the glued-word bug: the engine drops a text
    #     node's leading newline, which once turned "…menu.\n21 seeds" into
    #     "…menu.21 seeds" and "…DNUI\n*This catalogue" into "…DNUIThis".
    for risk in stats["glue_risks"]:
        problems.append(f"would glue in the handout: {risk}")
    for glued in re.findall(r"\bmenu\.[^\s]", blob):
        problems.append(f"glued text survived: {glued!r}")

    # (f) the section headings are real Word headings and carry the intended
    #     formatting (the built-in Heading styles must not leak through)
    from docx import Document

    doc = Document(str(docx_path))
    heading_paras = [p for p in doc.paragraphs if p.style.name.startswith("Heading")]
    for para in heading_paras:
        for run in para.runs:
            if run.font.italic or (run.font.underline is not None and run.font.underline):
                problems.append(f"heading has stray italic/underline: {para.text[:40]!r}")
                break
    for n in range(0, 8):
        if not any(h.startswith(f"{n}.") for h in heads):
            problems.append(f"section {n} heading not found in the outline")

    return problems


def main() -> int:
    if not MD_PATH.exists():
        raise SystemExit(f"source not found: {MD_PATH}")

    md_text = MD_PATH.read_text(encoding="utf-8")
    out_html, stats = md_to_html(md_text)

    print(f"source : {MD_PATH.relative_to(ROOT)}  ({len(md_text.splitlines())} lines)")
    print(f"  tables kept          : {stats['tables']}")
    print(f"  empty headers dropped: {stats['empty_headers_dropped']}")
    print(f"  links made absolute  : {stats['links_rewritten']}")
    print(f"  soft breaks joined   : {stats['soft_breaks_joined']}")
    print(f"  blockquotes inlined  : {stats['blockquotes_inlined']}")

    with tempfile.TemporaryDirectory() as tmp:
        html_path = Path(tmp) / "topic-catalogue.html"
        html_path.write_text(out_html, encoding="utf-8")
        print(f"target : {OUT_PATH.relative_to(ROOT)}")
        run_engine(html_path, OUT_PATH)

    print(f"  headings normalised  : {normalize_headings(OUT_PATH)}")

    problems = verify(md_text, OUT_PATH, stats)
    size_kb = OUT_PATH.stat().st_size / 1024
    if problems:
        print(f"\nFAILED verification ({len(problems)} problem(s)):")
        for p in problems:
            print(f"  - {p}")
        return 1

    _, heads, tables = docx_text(OUT_PATH)
    print(f"\nOK  {OUT_PATH.name}  {size_kb:.0f} KB")
    print(f"    {len(heads)} headings · {len(tables)} tables · "
          f"{len(SEED_IDS)} seed cards verified")
    return 0


if __name__ == "__main__":
    sys.exit(main())
