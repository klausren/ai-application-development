"""Deck validator for the AI Application Development course.

Usage:
    python scripts/validate_deck.py lectures/week-01/lecture-en.pptx [more.pptx ...]

Reports slide count and, for files whose name ends in -en.pptx, any slide that
still contains CJK characters (a sign that a zh/en field was swapped).
"""

import re
import sys

from pptx import Presentation

CJK = re.compile(r"[\u4e00-\u9fff]")


def slide_text(slide):
    parts = []
    for sh in slide.shapes:
        if sh.has_text_frame:
            parts.append(sh.text_frame.text)
        if sh.has_table:
            for row in sh.table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
    return "\n".join(parts)


def check(path):
    prs = Presentation(path)
    n = len(prs.slides)
    leaks = []
    if path.endswith("-en.pptx"):
        for i, s in enumerate(prs.slides, 1):
            if CJK.search(slide_text(s)):
                leaks.append(i)
    ok = "OK" if not leaks else "LEAK"
    print(f"{ok:4}  {path}  slides={n}  cjk_leak={leaks if leaks else 'none'}")
    return not leaks


if __name__ == "__main__":
    files = sys.argv[1:]
    if not files:
        print(__doc__)
        sys.exit(1)
    good = all(check(f) for f in files)
    sys.exit(0 if good else 2)
