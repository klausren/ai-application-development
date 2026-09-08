"""Deck builder for the AI Application Development course.

Renders a 16:9 (33.867 x 19.05 cm) lecture deck from a language-neutral slide
spec.  Every visual is a native PowerPoint shape, table or chart — no raster or
SVG assets are embedded, so PowerPoint never shows "picture can't be displayed".

Usage:
    python scripts/build_deck.py week-03
"""

import sys
from pptx import Presentation
from pptx.util import Cm, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION

# ---------------------------------------------------------------- palette ----
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
BG         = RGBColor(0xFF, 0xFF, 0xFF)
TEXT       = RGBColor(0x1F, 0x29, 0x37)
MUTED      = RGBColor(0x6B, 0x72, 0x80)
BORDER     = RGBColor(0xE5, 0xE7, 0xEB)
CARD       = RGBColor(0xF0, 0xF5, 0xFC)
ACCENT     = RGBColor(0x1E, 0x4F, 0xA8)
ACCENT_BR  = RGBColor(0x3D, 0x7B, 0xD9)
DEEP       = RGBColor(0x0E, 0x3F, 0x8C)
GREEN      = RGBColor(0x05, 0x96, 0x69)
GREEN_BG   = RGBColor(0xEC, 0xFD, 0xF5)
ORANGE     = RGBColor(0xD9, 0x77, 0x06)
ORANGE_BG  = RGBColor(0xFF, 0xFB, 0xEB)
RED        = RGBColor(0xB9, 0x1C, 0x1C)
RED_BG     = RGBColor(0xFE, 0xF2, 0xF2)
CODE_BG    = RGBColor(0x0F, 0x17, 0x2A)
CODE_TEXT  = RGBColor(0xE2, 0xE8, 0xF0)
GHOST      = RGBColor(0xF5, 0xF7, 0xFA)

FONT = "Source Han Sans SC"
CODE_FONT = "Menlo"

SW, SH = 33.867, 19.05
M = 2.22                      # outer margin
CW = SW - 2 * M               # content width
HEADER_H = 2.12
BODY_TOP = 2.90
FOOTER_Y = 18.04


# ---------------------------------------------------------------- helpers ----
def L(item, lang):
    """Pick a language from a {'zh': .., 'en': ..} value, or return plain text."""
    if isinstance(item, dict):
        return item.get(lang) or item.get("en") or item.get("zh") or ""
    return item


def _txbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP, wrap=True):
    tb = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tb, tf


def text(slide, x, y, w, h, runs, size=13, bold=False, color=TEXT,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.25,
         font=FONT, space_after=0):
    """runs: str  ->  single paragraph;  list[str] -> one paragraph per item."""
    tb, tf = _txbox(slide, x, y, w, h, anchor)
    paras = [runs] if isinstance(runs, str) else list(runs)
    for i, item in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        if space_after:
            p.space_after = Pt(space_after)
        if (isinstance(item, tuple) and len(item) == 2 and isinstance(item[1], dict)):
            items = [item]                      # (text, opts) -> one styled run
        elif isinstance(item, (list, tuple)):
            items = list(item)
        else:
            items = [item]
        for j, r in enumerate(items):
            txt, opts = (r if isinstance(r, tuple) else (r, {}))
            run = p.add_run()
            run.text = txt
            f = run.font
            f.name = opts.get("font", font)
            f.size = Pt(opts.get("size", size))
            f.bold = opts.get("bold", bold)
            f.italic = opts.get("italic", False)
            f.color.rgb = opts.get("color", color)
        if tf.word_wrap is None:
            tf.word_wrap = True
    return tb


def rect(slide, x, y, w, h, fill=WHITE, line=None, line_w=0.75,
         shape=MSO_SHAPE.RECTANGLE, adj=None):
    sp = slide.shapes.add_shape(shape, Cm(x), Cm(y), Cm(w), Cm(h))
    sp.shadow.inherit = False
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(line_w)
    sp.text_frame.word_wrap = True
    if adj is not None:
        try:
            sp.adjustments[0] = adj
        except Exception:
            pass
    return sp


def arrow(slide, x, y, w, h, color=ACCENT_BR, shape=MSO_SHAPE.RIGHT_ARROW):
    sp = rect(slide, x, y, w, h, fill=color, shape=shape)
    return sp


def hexc(h):
    h = h.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def col(v, default=TEXT):
    if v is None:
        return default
    if isinstance(v, RGBColor):
        return v
    return hexc(v)


# ------------------------------------------------------------ slide frame ----
def new_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, SW, SH, fill=BG)
    return s


def header(slide, title, right, lang, subtitle=None):
    rect(slide, 0, 0, SW, HEADER_H, fill=WHITE)
    rect(slide, 0, HEADER_H, SW, 0.04, fill=BORDER)
    if subtitle:
        text(slide, M, 0.62, CW - 8, 0.72, L(title, lang), size=19, bold=True, color=DEEP)
        text(slide, M, 1.38, CW - 8, 0.5, L(subtitle, lang), size=11.5, color=MUTED)
    else:
        text(slide, M, 0.71, CW - 8, 0.72, L(title, lang), size=19, bold=True, color=DEEP)
    if right:
        text(slide, SW - M - 7.6, 0.82, 7.6, 0.5, L(right, lang), size=11,
             color=MUTED, align=PP_ALIGN.RIGHT)


def footer(slide, week, page, total):
    text(slide, M, FOOTER_Y, 12, 0.45,
         f"AI Application Development · Week {week}", size=10, color=MUTED)
    text(slide, SW - M - 3, FOOTER_Y, 3, 0.45, f"{page:02d} / {total}",
         size=10, color=MUTED, align=PP_ALIGN.RIGHT)


# ------------------------------------------------------------- renderers ----
def r_cover(slide, sp, ctx):
    lang, week = ctx["lang"], ctx["week"]
    rect(slide, 0, 0, SW, 0.37, fill=ACCENT)
    rect(slide, 0, 0.37, SW, 0.16, fill=ACCENT_BR)
    text(slide, SW / 2, 2.38, SW / 2 - M, 12.7, f"{week:02d}", size=190,
         bold=True, color=CARD, align=PP_ALIGN.RIGHT)
    text(slide, M, 2.38, 20, 0.6, L(sp["kicker"], lang), size=11.5, bold=True,
         color=ACCENT)
    rect(slide, M, 3.52, 1.69, 0.11, fill=ACCENT)
    text(slide, M, 4.47, CW, 1.5, L(sp["title"], lang), size=40, bold=True, color=DEEP)
    if sp.get("title_sub"):
        text(slide, M, 6.2, CW, 2.7, L(sp["title_sub"], lang), size=25, color=TEXT)
    text(slide, M, 9.29, CW, 0.8, L(sp["subtitle"], lang), size=14, color=MUTED)
    chips = sp["chips"]
    gap = 0.62
    total_w = CW - gap * (len(chips) - 1)
    widths = [total_w * c[1] / sum(c[1] for c in chips) for c in chips]
    x = M
    for (label, _), w in zip(chips, widths):
        rect(slide, x, 14.71, w, 0.9, fill=CARD)
        text(slide, x + 0.48, 14.9, w - 0.96, 0.53, L(label, lang), size=11.5,
             bold=True, color=DEEP)
        x += w + gap
    rect(slide, M, 16.30, 1.16, 1.16, fill=DEEP)
    text(slide, M + 0.16, 16.56, 0.85, 0.64, f"W{week}", size=11, bold=True, color=WHITE)
    text(slide, M + 1.48, 16.62, 16, 0.56, L(sp.get("lecturer",
         "Lecturer: ____________  ·  School of Software Engineering"), lang),
         size=11, color=MUTED)


def r_contents(slide, sp, ctx):
    lang, week, total = ctx["lang"], ctx["week"], ctx["total"]
    header(slide, {"zh": "目录", "en": "Contents"}, sp.get("right"), lang)
    footer(slide, week, sp["page"], total)
    rect(slide, M, BODY_TOP, 7.94, 13.76, fill=CARD)
    text(slide, M + 0.69, 6.01, 6.56, 1.08, L(sp["count"], lang), size=26, bold=True, color=DEEP)
    text(slide, M + 0.69, 7.30, 6.56, 0.53, L(sp["count_label"], lang), size=11.5, bold=True, color=MUTED)
    rect(slide, M + 0.69, 8.31, 6.56, 0.05, fill=ACCENT_BR)
    text(slide, M + 0.69, 8.78, 6.56, 2.49, L(sp["tagline"], lang), size=12.5, color=TEXT, spacing=1.4)
    rect(slide, M + 0.69, 11.69, 6.56, 1.91, fill=WHITE, line=BORDER)
    text(slide, M + 1.06, 11.96, 5.82, 1.38, L(sp["timing"], lang), size=11.5,
         color=MUTED, spacing=1.35)

    x0, w = 10.90, 20.74
    y = 4.58
    for i, part in enumerate(sp["parts"]):
        rect(slide, x0, y, w, 2.12, fill=WHITE, line=BORDER)
        rect(slide, x0 + 0.21, y + 0.34, 1.38, 1.38, fill=CARD)
        text(slide, x0 + 0.21, y + 0.62, 1.38, 0.64, f"{i + 1:02d}", size=17,
             bold=True, color=ACCENT, align=PP_ALIGN.CENTER)
        text(slide, x0 + 2.12, y + 0.39, 12, 0.77, L(part["title"], lang), size=16,
             bold=True, color=DEEP)
        if part.get("title_en"):
            text(slide, x0 + 2.12 + 0, y + 1.19, 12, 0.45,
                 L(part["title_en"], lang), size=11.5, color=ACCENT_BR, bold=True)
        text(slide, x0 + 2.12, y + 1.79 if not part.get("title_en") else y + 1.72,
             17.36, 0.45, L(part["desc"], lang), size=11.5, color=MUTED)
        # chevron instead of an icon asset
        arrow(slide, x0 + w - 0.95, y + 0.79, 0.5, 0.53, color=ACCENT_BR)
        y += 2.09


def r_section(slide, sp, ctx):
    lang, week, total = ctx["lang"], ctx["week"], ctx["total"]
    rect(slide, 0, 0, SW, 0.26, fill=ACCENT)
    text(slide, SW - M - 18.5, 1.06, 18.52, 15.69, sp["num"], size=170, bold=True,
         color=CARD, align=PP_ALIGN.RIGHT)
    rect(slide, M + 0.16, 2.91, 3.44, 3.44, fill=DEEP)
    text(slide, M + 1.06, 3.65, 1.67, 1.98, sp["num"], size=32, bold=True,
         color=WHITE, align=PP_ALIGN.CENTER)
    text(slide, M + 4.39, 3.07, 14, 1.4, L(sp["title"], lang), size=30, bold=True, color=DEEP)
    text(slide, M + 4.39, 4.63, 14, 0.85, L(sp["title_en"], lang), size=15, color=ACCENT_BR)
    if sp.get("lead"):
        text(slide, M, 7.6, 24, 4.5, L(sp["lead"], lang), size=15, color=TEXT, spacing=1.5)
    cards = sp.get("pillars", [])
    if cards:
        gap = 0.63
        w = (CW - gap * (len(cards) - 1)) / len(cards)
        x = M
        for c in cards:
            rect(slide, x, 13.89, w, 2.25, fill=CARD)
            rect(slide, x, 13.89, w, 0.07, fill=ACCENT_BR)
            text(slide, x + 0.58, 14.37, w - 1.16, 0.71, L(c["title"], lang),
                 size=15, bold=True, color=DEEP)
            text(slide, x + 0.58, 15.19, w - 1.16, 0.48, L(c["sub"], lang),
                 size=11, color=MUTED)
            x += w + gap
    footer(slide, week, sp["page"], total)


def r_cards(slide, sp, ctx):
    lang, week, total = ctx["lang"], ctx["week"], ctx["total"]
    header(slide, sp["title"], sp.get("right"), lang, sp.get("subtitle"))
    footer(slide, week, sp["page"], total)
    cards = sp["cards"]
    cols = sp.get("cols", 2)
    gap = sp.get("gap", 0.6)
    top = sp.get("top", BODY_TOP + 0.35)
    h = sp.get("h", 4.2)
    w = (CW - gap * (cols - 1)) / cols
    for i, c in enumerate(cards):
        r, cidx = divmod(i, cols)
        x = M + cidx * (w + gap)
        y = top + r * (h + gap)
        fill = col(c.get("bg"), CARD)
        rect(slide, x, y, w, h, fill=fill, line=BORDER)
        rect(slide, x, y, 0.09, h, fill=col(c.get("accent"), ACCENT_BR))
        text(slide, x + 0.62, y + 0.52, w - 1.24, 0.62, L(c["title"], lang),
             size=15.5, bold=True, color=col(c.get("title_color"), DEEP))
        if c.get("title_en"):
            text(slide, x + 0.62, y + 1.16, w - 1.24, 0.42, L(c["title_en"], lang),
                 size=11, color=ACCENT_BR, bold=True)
        body_top = y + (1.72 if c.get("title_en") else 1.28)
        body = c.get("body") or []
        if isinstance(body, str):
            body = [body]
        text(slide, x + 0.62, body_top, w - 1.24, h - (body_top - y) - 0.4,
             [("• " + L(b, lang), {"size": c.get("body_size", 12)}) for b in body],
             size=c.get("body_size", 12), color=TEXT, spacing=1.4, space_after=4)
        if c.get("tag"):
            rect(slide, x + w - 3.4, y + h - 1.15, 2.8, 0.62, fill=col(c.get("tag_bg"), DEEP))
            text(slide, x + w - 3.4, y + h - 1.02, 2.8, 0.42, L(c["tag"], lang),
                 size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)


def r_table(slide, sp, ctx):
    lang, week, total = ctx["lang"], ctx["week"], ctx["total"]
    header(slide, sp["title"], sp.get("right"), lang, sp.get("subtitle"))
    footer(slide, week, sp["page"], total)
    cols = sp["cols"]
    rows = sp["rows"]
    top = sp.get("top", BODY_TOP + 0.3)
    w = sp.get("w", CW)
    x = sp.get("x", M)
    h = sp.get("h", min(13.0, 1.15 * (len(rows) + 1)))
    widths = sp.get("widths") or [w / len(cols)] * len(cols)
    shape = slide.shapes.add_table(len(rows) + 1, len(cols), Cm(x), Cm(top),
                                   Cm(w), Cm(h))
    tbl = shape.table
    for i, cwv in enumerate(widths):
        tbl.columns[i].width = Cm(cwv)
    for j, c in enumerate(cols):
        cell = tbl.cell(0, j)
        cell.text = L(c, lang)
        cell.fill.solid()
        cell.fill.fore_color.rgb = DEEP
        p = cell.text_frame.paragraphs[0]
        p.runs[0].font.size = Pt(sp.get("head_size", 12.5))
        p.runs[0].font.bold = True
        p.runs[0].font.name = FONT
        p.runs[0].font.color.rgb = WHITE
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.text = L(val, lang)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if i % 2 else CARD
            p = cell.text_frame.paragraphs[0]
            p.runs[0].font.size = Pt(sp.get("body_size", 11.5))
            p.runs[0].font.name = FONT
            p.runs[0].font.color.rgb = TEXT
            if j == 0 and sp.get("bold_first"):
                p.runs[0].font.bold = True
    if sp.get("note"):
        text(slide, x, top + h + 0.4, w, 1.2, L(sp["note"], lang), size=11.5,
             color=MUTED, spacing=1.35)


def r_flow(slide, sp, ctx):
    """Horizontal steps connected by arrows; optional note below."""
    lang, week, total = ctx["lang"], ctx["week"], ctx["total"]
    header(slide, sp["title"], sp.get("right"), lang, sp.get("subtitle"))
    footer(slide, week, sp["page"], total)
    steps = sp["steps"]
    gap = 0.95
    n = len(steps)
    aw = 0.85                                   # arrow width
    w = (CW - gap * (n - 1) - aw * (n - 1)) / n
    y = sp.get("y", 6.4)
    h = sp.get("h", 4.6)
    x = M
    for i, s in enumerate(steps):
        rect(slide, x, y, w, h, fill=col(s.get("bg"), CARD), line=BORDER)
        rect(slide, x, y, w, 0.1, fill=col(s.get("accent"), ACCENT_BR))
        rect(slide, x + w / 2 - 0.62, y + 0.55, 1.24, 1.24, fill=col(s.get("accent"), ACCENT_BR))
        text(slide, x + w / 2 - 0.62, y + 0.78, 1.24, 0.8, f"{i + 1}", size=18,
             bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        text(slide, x + 0.55, y + 2.05, w - 1.1, 0.62, L(s["title"], lang),
             size=14.5, bold=True, color=DEEP, align=PP_ALIGN.CENTER)
        if s.get("title_en"):
            text(slide, x + 0.55, y + 2.68, w - 1.1, 0.4, L(s["title_en"], lang),
                 size=10.5, color=ACCENT_BR, align=PP_ALIGN.CENTER)
        text(slide, x + 0.55, y + 3.2, w - 1.1, h - 3.4, L(s["desc"], lang),
             size=11.5, color=MUTED, align=PP_ALIGN.CENTER, spacing=1.35)
        x += w
        if i < n - 1:
            arrow(slide, x + (gap - aw) / 2, y + h / 2 - 0.32, aw, 0.64,
                  color=col(s.get("accent"), ACCENT_BR))
            x += gap
    if sp.get("note"):
        text(slide, M, y + h + 0.85, CW, 1.6, L(sp["note"], lang), size=12.5,
             color=TEXT, spacing=1.4)


def r_compare(slide, sp, ctx):
    """Two (or three) columns side by side with bullet lists."""
    lang, week, total = ctx["lang"], ctx["week"], ctx["total"]
    header(slide, sp["title"], sp.get("right"), lang, sp.get("subtitle"))
    footer(slide, week, sp["page"], total)
    cols = sp["columns"]
    gap = 0.85
    w = (CW - gap * (len(cols) - 1)) / len(cols)
    top = sp.get("top", BODY_TOP + 0.35)
    h = sp.get("h", 11.4)
    x = M
    for c in cols:
        rect(slide, x, top, w, h, fill=col(c.get("bg"), CARD), line=BORDER)
        rect(slide, x, top, w, 1.5, fill=col(c.get("accent"), DEEP))
        text(slide, x + 0.6, top + 0.42, w - 1.2, 0.7, L(c["title"], lang),
             size=16, bold=True, color=WHITE)
        items = c.get("items", [])
        text(slide, x + 0.6, top + 2.0, w - 1.2, h - 2.4,
             [("• " + L(i, lang)) for i in items], size=12.5, color=TEXT,
             spacing=1.5, space_after=8)
        x += w + gap
    if sp.get("note"):
        text(slide, M, top + h + 0.5, CW, 1.6, L(sp["note"], lang), size=12.5,
             color=TEXT, spacing=1.4)


def r_splitbar(slide, sp, ctx):
    """Proportional horizontal segments (train / val / test ...)."""
    lang, week, total = ctx["lang"], ctx["week"], ctx["total"]
    header(slide, sp["title"], sp.get("right"), lang, sp.get("subtitle"))
    footer(slide, week, sp["page"], total)
    segs = sp["segments"]
    total_v = sum(s["value"] for s in segs)
    y, h = sp.get("y", 5.6), sp.get("h", 3.1)
    x = M
    for s in segs:
        w = CW * s["value"] / total_v
        rect(slide, x, y, w - 0.12, h, fill=col(s.get("color"), ACCENT))
        text(slide, x, y + 0.55, w - 0.12, 0.8, L(s["title"], lang), size=17,
             bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        text(slide, x, y + 1.42, w - 0.12, 0.6, s.get("pct", ""), size=13,
             bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        text(slide, x + 0.5, y + h + 0.55, w - 1.0, 2.6, L(s["desc"], lang),
             size=12, color=TEXT, spacing=1.4)
        x += w
    if sp.get("note"):
        text(slide, M, y + h + 3.5, CW, 2.0, L(sp["note"], lang), size=12.5,
             color=DEEP, spacing=1.4)


def r_kfold(slide, sp, ctx):
    lang, week, total = ctx["lang"], ctx["week"], ctx["total"]
    header(slide, sp["title"], sp.get("right"), lang, sp.get("subtitle"))
    footer(slide, week, sp["page"], total)
    k = sp.get("k", 5)
    rows = sp.get("rows", k)
    x, y = sp.get("x", M), sp.get("y", 4.4)
    w = sp.get("w", 18.5)
    h, gap = 1.32, 0.42
    seg_w = w / k
    for i in range(rows):
        label = f"Fold {i + 1}" if lang == "en" else f"第 {i + 1} 折"
        text(slide, x, y + i * (h + gap) + 0.28, 3.0, 0.6, label, size=12,
             bold=True, color=DEEP)
        for j in range(k):
            sx = x + 3.3 + j * seg_w
            is_val = j == i
            rect(slide, sx, y + i * (h + gap), seg_w - 0.14, h,
                 fill=ORANGE if is_val else col(sp.get("train_color"), DEEP))
            txt = "val" if is_val else "train"
            if lang == "zh":
                txt = "验证" if is_val else "训练"
            text(slide, sx, y + i * (h + gap) + 0.36, seg_w - 0.14, 0.6, txt,
                 size=10.5, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    notes = sp.get("notes", [])
    if notes:
        text(slide, x + w + 1.1, y + 0.2, SW - M - (x + w + 1.1), 9.5,
             [("• " + L(n, lang)) for n in notes], size=12.5, color=TEXT,
             spacing=1.5, space_after=9)
    if sp.get("note"):
        text(slide, M, y + rows * (h + gap) + 0.7, CW, 1.8, L(sp["note"], lang),
             size=12.5, color=MUTED, spacing=1.4)


def r_chart(slide, sp, ctx):
    lang, week, total = ctx["lang"], ctx["week"], ctx["total"]
    header(slide, sp["title"], sp.get("right"), lang, sp.get("subtitle"))
    footer(slide, week, sp["page"], total)
    data = CategoryChartData()
    data.categories = sp["categories"]
    for s in sp["series"]:
        data.add_series(L(s["name"], lang), s["values"])
    x, y = sp.get("x", M), sp.get("y", BODY_TOP + 0.3)
    w, h = sp.get("w", 19.0), sp.get("h", 11.6)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS if sp.get("markers", True)
                                else XL_CHART_TYPE.LINE,
                                Cm(x), Cm(y), Cm(w), Cm(h), data)
    ch = gf.chart
    ch.has_title = False
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    ch.font.size = Pt(11)
    ch.font.name = FONT
    for i, (s, ser) in enumerate(zip(sp["series"], ch.plots[0].series)):
        c = col(s.get("color"), [ACCENT, ORANGE, GREEN, RED][i % 4])
        ser.format.line.color.rgb = c
        ser.format.line.width = Pt(2.5 if s.get("emphasis") else 1.75)
        ser.smooth = s.get("smooth", True)
        if s.get("no_markers"):
            ser.marker.format.fill.background()
    if sp.get("x_label"):
        ch.category_axis.axis_title.text_frame.text = L(sp["x_label"], lang)
        ch.category_axis.has_major_gridlines = False
    if sp.get("y_label"):
        ch.value_axis.axis_title.text_frame.text = L(sp["y_label"], lang)
    if sp.get("y_max") is not None:
        ch.value_axis.maximum_scale = sp["y_max"]
    notes = sp.get("notes", [])
    if notes:
        text(slide, x + w + 0.9, y + 0.3, SW - M - (x + w + 0.9), 9.0,
             [("• " + L(n, lang)) for n in notes], size=12.5, color=TEXT,
             spacing=1.45, space_after=8)
    if sp.get("note"):
        text(slide, M, y + h + 0.45, CW, 1.6, L(sp["note"], lang), size=12,
             color=MUTED, spacing=1.35)


def r_code(slide, sp, ctx):
    lang, week, total = ctx["lang"], ctx["week"], ctx["total"]
    header(slide, sp["title"], sp.get("right"), lang, sp.get("subtitle"))
    footer(slide, week, sp["page"], total)
    blocks = sp["blocks"]
    gap = 0.62
    n = len(blocks)
    w = (CW - gap * (n - 1)) / n
    top = sp.get("top", BODY_TOP + 0.35)
    h = sp.get("h", 11.3)
    x = M
    for b in blocks:
        rect(slide, x, top, w, h, fill=CODE_BG)
        rect(slide, x, top, w, 0.85, fill=RGBColor(0x1B, 0x26, 0x3D))
        text(slide, x + 0.45, top + 0.24, w - 3, 0.45, L(b["title"], lang),
             size=11.5, bold=True, color=RGBColor(0x93, 0xC5, 0xFD))
        code = b["code"]
        if isinstance(code, dict):                 # per-language code block
            code = code.get(lang) or code.get("en")
        lines = code
        if isinstance(lines, str):
            lines = lines.split("\n")
        text(slide, x + 0.55, top + 1.15, w - 1.1, h - 1.5,
             [(ln, {"font": CODE_FONT, "size": b.get("size", 11)}) for ln in lines],
             size=b.get("size", 11), color=CODE_TEXT, spacing=1.28)
        x += w + gap
    if sp.get("note"):
        text(slide, M, top + h + 0.5, CW, 1.7, L(sp["note"], lang), size=12,
             color=TEXT, spacing=1.35)


def r_checklist(slide, sp, ctx):
    lang, week, total = ctx["lang"], ctx["week"], ctx["total"]
    header(slide, sp["title"], sp.get("right"), lang, sp.get("subtitle"))
    footer(slide, week, sp["page"], total)
    items = sp["items"]
    top = sp.get("top", BODY_TOP + 0.3)
    h, gap = sp.get("h", 2.05), 0.42
    for i, it in enumerate(items):
        y = top + i * (h + gap)
        rect(slide, M, y, CW, h, fill=WHITE, line=BORDER)
        rect(slide, M, y, 0.09, h, fill=col(it.get("accent"), ACCENT_BR))
        if it.get("badge"):
            rect(slide, M + 0.5, y + 0.52, 2.5, 0.9, fill=col(it.get("accent"), ACCENT_BR))
            text(slide, M + 0.5, y + 0.72, 2.5, 0.5, L(it["badge"], lang), size=11.5,
                 bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        text(slide, M + 3.4, y + 0.32, CW - 4.0, 0.55, L(it["title"], lang),
             size=14, bold=True, color=DEEP)
        text(slide, M + 3.4, y + 0.95, CW - 4.0, h - 1.1, L(it["desc"], lang),
             size=11.5, color=MUTED, spacing=1.35)


def r_quiz(slide, sp, ctx):
    lang, week, total = ctx["lang"], ctx["week"], ctx["total"]
    header(slide, sp["title"], sp.get("right"), lang, sp.get("subtitle"))
    footer(slide, week, sp["page"], total)
    qs = sp["questions"]
    cols = sp.get("cols", 2)
    gap = 0.55
    w = (CW - gap * (cols - 1)) / cols
    top = sp.get("top", BODY_TOP + 0.25)
    rows = (len(qs) + cols - 1) // cols
    h = sp.get("h", (14.4 - (rows - 1) * gap) / rows)
    for i, q in enumerate(qs):
        r, c = divmod(i, cols)
        x = M + c * (w + gap)
        y = top + r * (h + gap)
        rect(slide, x, y, w, h, fill=CARD, line=BORDER)
        rect(slide, x + 0.55, y + 0.45, 2.5, 0.62, fill=DEEP)
        text(slide, x + 0.55, y + 0.58, 2.5, 0.42, L(q["tag"], lang), size=10.5,
             bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        text(slide, x + 3.35, y + 0.45, w - 3.9, 1.5, L(q["q"], lang), size=12.5,
             bold=True, color=DEEP, spacing=1.3)
        opts = q.get("options", [])
        text(slide, x + 0.55, y + 2.15, w - 1.1, h - 2.4,
             [(L(o, lang), {"size": 11.5}) for o in opts], size=11.5, color=TEXT,
             spacing=1.4, space_after=3)


def r_summary(slide, sp, ctx):
    lang, week, total = ctx["lang"], ctx["week"], ctx["total"]
    header(slide, sp["title"], sp.get("right"), lang, sp.get("subtitle"))
    footer(slide, week, sp["page"], total)
    pts = sp["points"]
    w = sp.get("w", 18.6)
    top = sp.get("top", BODY_TOP + 0.3)
    h, gap = 1.85, 0.42
    for i, p in enumerate(pts):
        y = top + i * (h + gap)
        rect(slide, M, y, w, h, fill=CARD)
        rect(slide, M, y, 1.32, h, fill=DEEP)
        text(slide, M, y + 0.55, 1.32, 0.7, f"{i + 1}", size=19, bold=True,
             color=WHITE, align=PP_ALIGN.CENTER)
        text(slide, M + 1.75, y + 0.35, w - 2.4, h - 0.7, L(p, lang), size=13,
             color=TEXT, anchor=MSO_ANCHOR.MIDDLE, spacing=1.35)
    todos = sp.get("todos", [])
    if todos:
        x = M + w + 0.85
        tw = SW - M - x
        rect(slide, x, top, tw, 1.35, fill=DEEP)
        text(slide, x + 0.6, top + 0.36, tw - 1.2, 0.6,
             L(sp.get("todo_title", {"zh": "课后待办", "en": "After Class"}), lang),
             size=14, bold=True, color=WHITE)
        for i, t in enumerate(todos):
            y = top + 1.85 + i * 2.95
            rect(slide, x, y, tw, 2.6, fill=WHITE, line=BORDER)
            text(slide, x + 0.6, y + 0.42, tw - 1.2, 0.6, L(t["title"], lang),
                 size=13.5, bold=True, color=DEEP)
            text(slide, x + 0.6, y + 1.12, tw - 1.2, 1.3, L(t["desc"], lang),
                 size=11.5, color=MUTED, spacing=1.35)


def r_closing(slide, sp, ctx):
    lang, week = ctx["lang"], ctx["week"]
    rect(slide, 0, 0, SW, 0.26, fill=ACCENT)
    text(slide, SW - M - 18.5, 1.06, 18.52, 15.69, f"{week:02d}", size=170,
         bold=True, color=CARD, align=PP_ALIGN.RIGHT)
    rect(slide, M + 0.16, 2.6, 1.16, 1.16, fill=DEEP)
    text(slide, M + 0.32, 2.86, 0.85, 0.64, f"W{week}", size=11, bold=True, color=WHITE)
    rect(slide, M, 4.4, 1.69, 0.11, fill=ACCENT)
    text(slide, M, 5.3, 24, 5.5, L(sp["quote"], lang), size=27, bold=True,
         color=DEEP, spacing=1.35)
    if sp.get("quote_sub"):
        text(slide, M, 9.4, 24, 1.2, L(sp["quote_sub"], lang), size=13.5, color=MUTED)
    text(slide, M, 11.6, 22, 0.9, L(sp.get("thanks", {"zh": "谢谢大家", "en": "Thank you"}), lang),
         size=22, bold=True, color=TEXT)
    rect(slide, M, 13.3, CW, 2.0, fill=CARD, line=BORDER)
    text(slide, M + 0.7, 13.62, CW - 1.4, 0.5,
         L(sp.get("next_label", {"zh": "下周预告", "en": "Next week"}), lang),
         size=11.5, bold=True, color=ACCENT)
    text(slide, M + 0.7, 14.15, CW - 1.4, 0.9, L(sp["next"], lang), size=14,
         color=DEEP, spacing=1.3)
    text(slide, M, FOOTER_Y, 20, 0.45, L(sp.get(
        "sign", f"AI Application Development · Week {week} Lecture · {{n}} / {{n}}"), lang)
        .replace("{n}", str(ctx["total"])), size=10, color=MUTED)


RENDERERS = {
    "cover": r_cover, "contents": r_contents, "section": r_section,
    "cards": r_cards, "table": r_table, "flow": r_flow, "compare": r_compare,
    "splitbar": r_splitbar, "kfold": r_kfold, "chart": r_chart, "code": r_code,
    "checklist": r_checklist, "quiz": r_quiz, "summary": r_summary,
    "closing": r_closing,
}


def build(spec, out_path, lang):
    prs = Presentation()
    prs.slide_width = Cm(SW)
    prs.slide_height = Cm(SH)
    slides = spec["slides"]
    ctx = {"lang": lang, "week": spec["week"], "total": len(slides)}
    for i, sp in enumerate(slides, start=1):
        sp = dict(sp)
        sp["page"] = i
        slide = new_slide(prs)
        RENDERERS[sp["type"]](slide, sp, ctx)
    prs.save(out_path)
    return len(slides)


if __name__ == "__main__":
    import importlib
    week = sys.argv[1] if len(sys.argv) > 1 else "week-03"
    mod = importlib.import_module(f"deck_specs.{week.replace('-', '_')}")
    spec = mod.SPEC
    n = build(spec, f"/tmp/{week}-en.pptx", "en")
    n = build(spec, f"/tmp/{week}-zh.pptx", "zh")
    print("slides:", n)
